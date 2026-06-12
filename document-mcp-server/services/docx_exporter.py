import re
from pathlib import Path

from docx import Document


def export_report_docx(topic: str, artifact_type: str, title: str, content: str, output_dir: str) -> dict:
    if not content or not content.strip():
        raise ValueError("报告内容不能为空，无法导出 Word 文件")

    export_dir = Path(output_dir).resolve()
    export_dir.mkdir(parents=True, exist_ok=True)

    file_name = build_file_name(topic, artifact_type)
    file_path = export_dir / file_name

    document = Document()
    document.add_heading(title or "课设报告导出", level=0)
    document.add_paragraph(f"课设题目：{topic or '未命名课设'}")
    document.add_paragraph("")

    for block in split_blocks(content):
        if block.startswith("# "):
            document.add_heading(block[2:].strip(), level=1)
        elif block.startswith("## "):
            document.add_heading(block[3:].strip(), level=2)
        elif block.startswith("### "):
            document.add_heading(block[4:].strip(), level=3)
        else:
            document.add_paragraph(block)

    document.save(file_path)

    return {
        "file_name": file_name,
        "file_path": str(file_path),
    }


def build_file_name(topic: str, artifact_type: str) -> str:
    normalized_topic = sanitize_file_name(topic or "课设报告")
    normalized_type = sanitize_file_name(artifact_type or "report")
    return f"{normalized_topic}-{normalized_type}.docx"


def sanitize_file_name(value: str) -> str:
    sanitized = re.sub(r"[\\\\/:*?\"<>|]+", "-", value).strip()
    sanitized = re.sub(r"\s+", "-", sanitized)
    return sanitized[:80] or "课设报告"


def split_blocks(content: str) -> list[str]:
    blocks = []
    for block in content.replace("\r\n", "\n").split("\n\n"):
        normalized = block.strip()
        if normalized:
            blocks.append(normalized)
    return blocks
